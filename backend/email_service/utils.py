import re
import csv
import io
from typing import List, Dict, Tuple


def replace_template_variables(text: str, row_data: Dict) -> str:
    """
    Replace template variables in text with values from row_data.
    
    Example: "Hello {name}, your email is {email}" with {"name": "John", "email": "john@example.com"}
    returns: "Hello John, your email is john@example.com"
    
    Args:
        text: Text containing template variables like {column_name}
        row_data: Dictionary with column names as keys and values
        
    Returns:
        Text with all template variables replaced
    """
    def replacer(match):
        variable_name = match.group(1).strip()
        # Return the value if exists, otherwise return the original placeholder
        return str(row_data.get(variable_name, match.group(0)))
    
    # Find all {variable_name} patterns and replace them
    return re.sub(r'\{([^}]+)\}', replacer, text)


def parse_csv_file(file_content: bytes) -> Tuple[List[Dict], List[str]]:
    """
    Parse CSV file and return list of row dictionaries.
    
    Args:
        file_content: Binary content of CSV file
        
    Returns:
        Tuple of (list of dicts with column data, list of column names)
    """
    try:
        # Decode the file content
        text = file_content.decode('utf-8')
        
        # Parse CSV
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
        
        if not rows:
            return [], []
        
        column_names = list(rows[0].keys())
        return rows, column_names
    
    except Exception as e:
        raise ValueError(f"Error parsing CSV file: {str(e)}")


def parse_excel_file(file_content: bytes) -> Tuple[List[Dict], List[str]]:
    """
    Parse Excel file and return list of row dictionaries.
    
    Args:
        file_content: Binary content of Excel file (.xlsx)
        
    Returns:
        Tuple of (list of dicts with column data, list of column names)
    """
    try:
        import openpyxl
        from io import BytesIO
        
        # Load workbook from bytes
        wb = openpyxl.load_workbook(BytesIO(file_content))
        ws = wb.active
        
        # Get header row (first row)
        headers = [cell.value for cell in ws[1]]
        
        # Convert to dictionary format
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for idx, header in enumerate(headers):
                if header is not None:
                    row_dict[str(header).strip()] = str(row[idx]) if row[idx] is not None else ""
            rows.append(row_dict)
        
        return rows, [str(h).strip() for h in headers if h is not None]
    
    except ImportError:
        raise ValueError("openpyxl is required for Excel file support. Install it with: pip install openpyxl")
    except Exception as e:
        raise ValueError(f"Error parsing Excel file: {str(e)}")
